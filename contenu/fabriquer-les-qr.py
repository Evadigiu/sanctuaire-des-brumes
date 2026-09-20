# -*- coding: utf-8 -*-
"""Fabrique les QR codes des bornes, a partir du plan du parcours.

    python3 contenu/fabriquer-les-qr.py

Produit deux choses, dans qr-codes/ :

  * un PNG par borne, SANS aucun texte. C'est l'image a integrer dans
    l'affichette. Le nom du fichier sert a s'y retrouver, il n'apparait
    nulle part sur l'image.

  * planche-de-pose.html, a imprimer pour l'equipe qui installe : la, tout
    est ecrit, le lieu, la borne, la position dans chaque sens.

Cette separation n'est pas cosmetique. Ce qui est visible dans le parc ne
doit rien raconter : une affichette portant « Indice : les jumelles, la
carcasse » spoilerait l'enigme a tous ceux qui passent devant.
"""
import os, sys, json, base64, io as _io, html
try:
    import qrcode
    from qrcode.constants import ERROR_CORRECT_H
except ImportError:
    sys.exit("Il manque la bibliotheque qrcode : pip install \"qrcode[pil]\"")

ICI    = os.path.dirname(os.path.abspath(__file__))
RACINE = os.path.dirname(ICI)
SORTIE = os.path.join(RACINE, "qr-codes")

def slug(t):
    import unicodedata, re
    d = unicodedata.normalize("NFKD", t.lower())
    d = "".join(c for c in d if not unicodedata.combining(c))
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", d)).strip("-")[:40]

plan = json.loads(open(os.path.join(RACINE, "assets", "js", "parcours.js"),
                       encoding="utf-8").read()
                  .split("const PARCOURS =", 1)[1].rstrip().rstrip(";"))

# Le domaine vient du meme endroit que les adresses du document, pour qu'ils
# ne puissent pas diverger.
src = open(os.path.join(ICI, "fabriquer-les-pages.py"), encoding="utf-8").read()
DOMAINE = src.split('DOMAINE = "', 1)[1].split('"', 1)[0]

# Le lieu de chaque borne, lu dans le cahier.
from openpyxl import load_workbook
par = load_workbook(os.path.join(ICI, "sanctuaire-cahier-de-contenu.xlsx"),
                    data_only=True)["Parcours"]
lieux = {}
for r in par.iter_rows(min_row=3, values_only=True):
    if r[0]: lieux[str(r[0]).strip()] = (str(r[3]).strip() if r[3] else "")

os.makedirs(SORTIE, exist_ok=True)
for f in os.listdir(SORTIE):
    if f.endswith((".png", ".html")): os.remove(os.path.join(SORTIE, f))

secours = {c: n for n, c in plan["secours"].items()}
fiches = []

for code in plan["pages"]:
    num = secours[code]
    url = "%s/%s" % (DOMAINE, plan["pages"][code])

    # Correction d'erreur maximale : une affichette exposee dehors pendant
    # dix-sept jours prend la pluie, la poussiere et les doigts. A ce niveau,
    # le code reste lisible avec pres d'un tiers de sa surface abimee.
    qr = qrcode.QRCode(version=None, error_correction=ERROR_CORRECT_H,
                       box_size=20, border=4)
    qr.add_data(url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")

    nom = "borne-%s-%s.png" % (num, slug(plan["noms"][code]))
    img.save(os.path.join(SORTIE, nom))

    tampon = _io.BytesIO(); img.save(tampon, format="PNG")
    fiches.append({
        "code": code, "num": num, "url": url,
        "nom": plan["noms"][code], "lieu": lieux.get(code, ""),
        "A": plan["ordre"]["A"].get(code), "B": plan["ordre"]["B"].get(code),
        "fichier": nom,
        "b64": base64.b64encode(tampon.getvalue()).decode(),
        "px": img.size[0],
    })

fiches.sort(key=lambda f: (f["A"] is None, f["A"] or f["B"] or 99))

# ---------- La planche de pose, pour l'equipe ----------
e = lambda t: html.escape(str(t) if t is not None else "—")
lignes = "".join(
  '<section class="fiche">'
  '<img src="data:image/png;base64,%s" alt="">'
  '<div class="infos">'
  '<div class="num">%s</div>'
  '<div class="borne">%s</div>'
  '<div class="lieu">%s</div>'
  '<table><tr><th>Sens A</th><td>%s</td><th>Sens B</th><td>%s</td></tr></table>'
  '<div class="url">%s</div>'
  '<div class="fichier">%s</div>'
  '</div></section>' % (f["b64"], e(f["num"]), e(f["nom"]), e(f["lieu"] or "lieu à préciser"),
                        e(f["A"]), e(f["B"]), e(f["url"]), e(f["fichier"]))
  for f in fiches)

open(os.path.join(SORTIE, "planche-de-pose.html"), "w", encoding="utf-8").write("""<!DOCTYPE html>
<html lang="fr"><head><meta charset="UTF-8">
<title>Planche de pose — Le Sanctuaire des Brumes</title>
<style>
  body { font-family: -apple-system, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
         margin: 24px; color: #1a1a1a; }
  h1 { font-size: 1.4rem; margin: 0 0 4px; }
  .chapo { color: #555; font-size: .9rem; max-width: 46rem; margin-bottom: 24px; }
  .chapo strong { color: #a33; }
  .fiche { display: flex; gap: 20px; align-items: center; border: 1px solid #ddd;
           border-radius: 8px; padding: 16px; margin-bottom: 14px;
           break-inside: avoid; page-break-inside: avoid; }
  .fiche img { width: 150px; height: 150px; flex: none; }
  .num { font-size: 2.4rem; font-weight: 700; letter-spacing: 2px; }
  .borne { font-size: 1.05rem; font-weight: 600; margin-top: 2px; }
  .lieu { color: #555; margin-bottom: 8px; }
  table { border-collapse: collapse; font-size: .85rem; margin-bottom: 8px; }
  th, td { border: 1px solid #ddd; padding: 3px 10px; text-align: left; }
  th { background: #f4f4f4; font-weight: 600; }
  .url { font-family: ui-monospace, Menlo, Consolas, monospace; font-size: .78rem; color: #666; }
  .fichier { font-size: .75rem; color: #999; margin-top: 4px; }
  @media print { body { margin: 10mm; } .chapo { font-size: .8rem; } }
</style></head><body>
<h1>Planche de pose, Le Sanctuaire des Brumes</h1>
<p class="chapo">Un QR par borne, dans l'ordre du sens A. Le nombre à quatre chiffres
doit figurer sur l'affichette à côté du code : c'est ce que le joueur saisira si le
QR refuse de se lire.<br>
<strong>Ce document est réservé à l'équipe.</strong> Le nom de la borne ne doit jamais
apparaître sur une affichette posée dans le parc : « Indice : les jumelles, la carcasse »
raconterait l'énigme à tous ceux qui passent devant. Sur l'affichette, le QR et le
nombre, rien d'autre.</p>
""" + lignes + "</body></html>")

print("%d QR codes dans qr-codes/" % len(fiches))
print("Planche de pose : qr-codes/planche-de-pose.html")
print("Domaine encode  : %s" % DOMAINE)
print("Taille des images : %d x %d pixels" % (fiches[0]["px"], fiches[0]["px"]))
